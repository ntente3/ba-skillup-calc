"""
Pull the results the workbook has already computed and keep them as an answer key.

The distributed workbook caches its last computed result in the cells. Comparing the
ported core against those, row for row, confirms that formulas transcribed by eye really
do produce the same values.

BD sheet (1. BD·노트·크레딧 계산), per student:
  S~V   BD consumed      basic/normal/advanced/superior
  W~AA  notes consumed   basic/normal/advanced/superior/secret
  AB    credits consumed
Gear sheet (3. 장비 계산):
  AW    gear credits
"""
import json, os, sys
import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from student_ids import load as load_reg, resolve


def _n(v):
    """Leave an empty cell as None.

    The workbook has no credit formula at all on collab student rows. Reading those as 0
    would amount to claiming "the workbook computed 0", and the check would then report a
    mismatch that does not exist.
    """
    return int(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else None


def _z(v):
    return _n(v) or 0


if __name__ == "__main__":
    wb = openpyxl.load_workbook(sys.argv[1], data_only=True)
    bd, eq = wb["1. BD·노트·크레딧 계산"], wb["3. 장비 계산"]
    reg = load_reg()

    brows = [r for r in range(19, 371) if bd.cell(r, 1).value is not None]
    erows = [r for r in range(13, 285) if eq.cell(r, 1).value is not None]

    out = []
    for rb, re_ in zip(brows, erows):
        name = bd.cell(rb, 2).value
        sid = resolve(reg, name)
        out.append({
            "sid": sid,
            "name": name,
            "own": bd.cell(rb, 1).value is True,
            "bd":     [_z(bd.cell(rb, c).value) for c in range(19, 23)],   # S~V
            "note":   [_z(bd.cell(rb, c).value) for c in range(23, 28)],   # W~AA
            "credit": _n(bd.cell(rb, 28).value),                            # AB
            "gearCredit": _n(eq.cell(re_, 49).value),                       # AW
        })
    wb.close()
    p = os.path.join(os.path.dirname(__file__), "..", "data", "expected.json")
    with open(p, "w", encoding="utf-8") as f:
        json.dump(out, f, ensure_ascii=False, separators=(",", ":"))
    nz = sum(1 for r in out if any(r["bd"]) or any(r["note"]) or r["credit"])
    blank = sum(1 for r in out if r["credit"] is None)
    print(f"rows with no credit formula (collab and the like): {blank} — left null, excluded from the check")
    print(f"answer key: {len(out)} rows ({nz} students with values) -> {p}")
