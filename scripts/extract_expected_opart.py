"""
Pull the cached results of the orb calculation sheet and keep them as an answer key.

This sheet is a dynamic array spill, so it has no inputs — it shows only those students
that actually need materials. Column K holds the student name (spilled), and from column T
onwards there are 20 materials x 4 grades = 80 columns.
"""
import json, os, sys
import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from student_ids import load as load_reg, resolve

FIRST_COL = 20   # T
N_MAT, N_GRADE = 20, 4


def _n(v):
    return int(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else 0


if __name__ == "__main__":
    wb = openpyxl.load_workbook(sys.argv[1], data_only=True)
    op = wb["2. 오파츠 계산"]
    reg = load_reg()

    out = []
    for r in range(12, 178):
        name = op.cell(r, 11).value
        if not isinstance(name, str) or name in ("#N/A", ""):
            continue
        sid = resolve(reg, name)
        if sid is None:
            continue
        grid = [
            [_n(op.cell(r, FIRST_COL + mi * N_GRADE + g).value) for g in range(N_GRADE)]
            for mi in range(N_MAT)
        ]
        out.append({"sid": sid, "name": name, "opart": grid})
    wb.close()

    p = os.path.join(os.path.dirname(__file__), "..", "data", "expected_opart.json")
    with open(p, "w", encoding="utf-8") as f:
        json.dump(out, f, ensure_ascii=False, separators=(",", ":"))
    nz = sum(1 for r in out if any(any(g) for g in r["opart"]))
    cells = sum(1 for r in out for g in r["opart"] for v in g if v)
    print(f"orb answer key: {len(out)} students ({nz} with values, {cells} non-zero cells) -> {p}")
